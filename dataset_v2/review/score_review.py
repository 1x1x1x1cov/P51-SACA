#!/usr/bin/env python3
"""Score the app against a filled-in clinician review sheet.

    python score_review.py clinician_review_sheet_FILLED.xlsx

Reads the 'Review' sheet (column A = ID, column C = rating) and review_answer_key.csv.
Reports agreement of the clinician with (a) the app and (b) the dataset's rule labels,
so you can see whether the labels or the model are the problem.
"""
import os
import sys

import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score, confusion_matrix

HERE = os.path.dirname(os.path.abspath(__file__))
LAB = ["LOW", "MEDIUM", "HIGH", "CRITICAL"]
R = {l: i for i, l in enumerate(LAB)}


def main(path):
    ws = load_workbook(path, data_only=True)["Review"]
    rows = []
    for r in range(5, ws.max_row + 1):
        i, rating, comment = ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        if i and rating:
            rating = str(rating).strip().upper()
            if rating not in R:
                sys.exit(f"row {r}: '{rating}' is not LOW/MEDIUM/HIGH/CRITICAL")
            rows.append((i, rating, comment or ""))
    if not rows:
        sys.exit("no ratings found in column C")
    d = pd.DataFrame(rows, columns=["id", "clinician", "comment"]).merge(
        pd.read_csv(os.path.join(HERE, "review_answer_key.csv")), on="id")
    out = [f"rated {len(d)} of the sampled sets"]
    for name, col in (("APP (LR + ceiling + red-flag floor)", "app_prediction"), ("dataset RULE LABELS", "rule_label")):
        c, m = d.clinician.map(R), d[col].map(R)
        out.append(f"\n== clinician vs {name}")
        out.append(f"exact agreement      {(c == m).mean():.1%}")
        out.append(f"within one tier      {((c - m).abs() <= 1).mean():.1%}")
        out.append(f"weighted kappa       {cohen_kappa_score(c, m, weights='quadratic'):.3f}")
        out.append(f"under-triage (model below clinician): {(m < c).sum()}   over-triage: {(m > c).sum()}")
        out.append(f"dangerous under-triage (clinician HIGH/CRITICAL, model LOW/MEDIUM): {((c >= 2) & (m <= 1)).sum()} of {(c >= 2).sum()}")
        out.append("confusion (rows = clinician, cols = model; LOW MED HIGH CRIT):")
        for row in confusion_matrix(c, m, labels=range(4)):
            out.append("   " + " ".join(f"{v:4d}" for v in row))
    big = d[(d.clinician.map(R) - d.app_prediction.map(R)).abs() >= 2]
    out.append(f"\n== app differs from clinician by 2+ tiers: {len(big)}")
    for r in big.itertuples():
        out.append(f"   {r.id}  clinician {r.clinician}  app {r.app_prediction}  rule {r.rule_label}  | {r.symptom_text}  | {r.comment}")
    text = "\n".join(out)
    print(text)
    open(os.path.join(HERE, "review_score.txt"), "w").write(text + "\n")


if __name__ == "__main__":
    main(sys.argv[1])
